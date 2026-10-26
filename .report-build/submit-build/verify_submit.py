from pathlib import Path
import json, re, os, hashlib, zipfile, csv
from urllib.parse import unquote
from docx import Document
from lxml import etree

root=Path(__file__).resolve().parents[2]
out=root/'Submit'
build=Path(__file__).resolve().parent
files=json.loads((build/'docx_outputs.json').read_text(encoding='utf-8'))
originals=[]
for folder in ['docs','scripts','images','汇报成果']:
    for src in (root/folder).rglob('*'):
        if src.is_file():
            target=out/'归档_历史材料_不放入提交压缩包'/src.relative_to(root)
            if not target.exists():
                target.parent.mkdir(parents=True,exist_ok=True);target.write_bytes(src.read_bytes())
            if target.read_bytes()!=src.read_bytes():
                target.write_bytes(src.read_bytes())
                print('UPDATED ARCHIVE',src.relative_to(root))
            originals.append(src)

# Include Drawio files and image exports added or updated while organizing.
for ext,folder in [('drawio','模型源文件_Drawio'),('png','模型截图')]:
    destdir=out/'作业2_需求建模成果'/'03-建模图'/folder
    for src in (root/'images').glob('*.'+ext):
        matches=[p for p in destdir.glob('*.'+ext) if p.name.endswith('-'+src.stem+'-9组-1011101011.'+ext)]
        dest=matches[0] if matches else destdir/f'{len(list(destdir.glob("*."+ext)))+1:02d}-{src.stem}-9组-1011101011.{ext}'
        if not dest.exists() or dest.read_bytes()!=src.read_bytes():dest.write_bytes(src.read_bytes())

if '--fix-links' in __import__('sys').argv:
    for path in files:
        path=Path(path)
        package=next(p for p in path.parents if p.name.startswith('作业'))
        source_dir=package/'文本源稿' if path.name.startswith(('01-需求调研报告','01-SRS','02-SPEC')) else path.parent
        d=Document(path)
        for rel in d.part.rels.values():
            if rel.reltype.endswith('/hyperlink') and rel.is_external:
                target=rel.target_ref
                if re.match(r'^(https?://|mailto:|#)',target):continue
                resolved=(source_dir/unquote(target)).resolve()
                assert resolved.is_file(),(path,target,resolved)
                rel._target=Path(os.path.relpath(resolved,path.parent)).as_posix()
        d.save(path)
    a=out/'作业1_需求调研成果';b=out/'作业2_需求建模成果'
    for n1,n2,title in [(4,6,'修改说明与版本记录'),(5,7,'AI使用声明'),(6,8,'提交材料自查清单')]:
        src=a/f'{n1:02d}-{title}-9组-1011101011.docx';dest=b/f'{n2:02d}-{title}-9组-1011101011.docx';dest.write_bytes(src.read_bytes())

checked_links=0
for path in out.rglob('*.docx'):
    if '归档_' in str(path) or '/原始调研材料/' in path.as_posix():continue
    d=Document(path)
    for rel in d.part.rels.values():
        if rel.reltype.endswith('/hyperlink') and rel.is_external and not re.match(r'^(https?://|mailto:|#)',rel.target_ref):
            assert (path.parent/unquote(rel.target_ref)).is_file(),(path,rel.target_ref)
            checked_links+=1
for package in [out/'作业1_需求调研成果',out/'作业2_需求建模成果']:
    for path in (package/'文本源稿').glob('*.md'):
        for target in re.findall(r'!?\[[^\]]*\]\(([^)]+)\)',path.read_text(encoding='utf-8')):
            if re.match(r'^(https?://|mailto:|#)',target):continue
            assert (path.parent/target).is_file(),(path,target)
            checked_links+=1

# Check all numbered requirements survive the conversion into Word.
def wordtext(f):
    d=Document(f)
    return '\n'.join([p.text for p in d.paragraphs]+[c.text for t in d.tables for row in t.rows for c in row.cells])
for source,label in [('系统需求规格说明书_SRS.md','01-SRS'),('核心模块开发规范_SPEC.md','02-SPEC')]:
    ids=set(re.findall(r'\b(?:FR|UC|AC|SPEC)-[A-Z]+-\d+\b',(root/'docs'/source).read_text(encoding='utf-8')))
    f=next((out/'作业2_需求建模成果').glob(label+'*.docx'))
    txt=wordtext(f)
    if label=='01-SRS':assert len(Document(f).inline_shapes)==2
    assert all(x in txt for x in ids),(source,ids-set(re.findall(r'\b(?:FR|UC|AC|SPEC)-[A-Z]+-\d+\b',txt)))

if '--pack' in __import__('sys').argv:
    for folder,suffix in [('作业1_需求调研成果','需求调研成果'),('作业2_需求建模成果','需求建模成果')]:
        package=out/folder
        dest=out/f'9组-1011101011-{suffix}.zip'
        with zipfile.ZipFile(dest,'w',zipfile.ZIP_DEFLATED) as z:
            for f in sorted(package.rglob('*')):
                if f.is_file():z.write(f,f.relative_to(package).as_posix())
        with zipfile.ZipFile(dest) as z:
            assert z.testzip() is None
            for f in package.rglob('*'):
                if f.is_file():assert z.read(f.relative_to(package).as_posix())==f.read_bytes()
        print('ZIP VERIFIED',dest.name,dest.stat().st_size)
    with (out/'文件清单.csv').open('w',encoding='utf-8-sig',newline='') as fp:
        writer=csv.writer(fp);writer.writerow(['文件相对路径','字节数','SHA256','用途'])
        for f in sorted(out.rglob('*')):
            if f.is_file() and f.name!='文件清单.csv':
                category='作业1' if '作业1_' in str(f) else '作业2' if '作业2_' in str(f) else '历史归档' if '归档_' in str(f) else '打包或说明'
                writer.writerow([f.relative_to(out).as_posix(),f.stat().st_size,hashlib.sha256(f.read_bytes()).hexdigest(),category])
    (build/'verification.json').write_text(json.dumps({'preserved_source_files':len(originals),'verified_local_links':checked_links,'docx_authored':len(files),'zip_integrity':'passed'},ensure_ascii=False,indent=2),encoding='utf-8')
print('CHECKED',len(originals),'preserved originals;',checked_links,'links')
