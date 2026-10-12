from pathlib import Path
import json
from docx import Document
from docx.oxml.ns import qn
build=Path(__file__).resolve().parent
for file in json.loads((build/'docx_outputs.json').read_text(encoding='utf-8')):
    path=Path(file);d=Document(path)
    for style in d.styles:
        if style.name in ['Title','Subtitle'] or style.name.startswith('Heading'):
            style.font.italic=False
            for border in style.element.xpath('./w:pPr/w:pBdr'):border.getparent().remove(border)
    for p in d.paragraphs:
        if p.style.name=='Title':
            for border in p._p.xpath('./w:pPr/w:pBdr'):border.getparent().remove(border)
    for table in d.tables:
        for row in list(table.rows)[1:]:
            if all(not c.text.strip() for c in row.cells):table._tbl.remove(row._tr)
        for p in table.rows[0].cells:
            for paragraph in p.paragraphs:paragraph.paragraph_format.keep_with_next=True
    d.save(path)
out=build.parents[1]/'Submit'
a=out/'作业1_需求调研成果';b=out/'作业2_需求建模成果'
for n1,n2,title in [(4,6,'修改说明与版本记录'),(5,7,'AI使用声明'),(6,8,'提交材料自查清单')]:
    src=a/f'{n1:02d}-{title}-组号待填-团队名待填.docx';dest=b/f'{n2:02d}-{title}-组号待填-团队名待填.docx';dest.write_bytes(src.read_bytes())
print('Polished 7 unique Word files and 3 matching copies')
