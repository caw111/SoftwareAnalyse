from pathlib import Path
import re, shutil, subprocess, json, hashlib, zipfile, os, sys
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'Submit'
BUILD=Path(__file__).resolve().parent
TAG='组号待填-团队名待填'
DATE='2026-10-12'
if OUT.exists() and '--resume' not in sys.argv:
    raise SystemExit('Submit already exists; refusing to overwrite')
OUT.mkdir(exist_ok=True)
A=OUT/'作业1_需求调研成果'
B=OUT/'作业2_需求建模成果'
ARCH=OUT/'归档_历史材料_不放入提交压缩包'
for p in [A,B,ARCH]:p.mkdir(exist_ok=True)

def named(n,title,ext):return f'{n:02d}-{title}-{TAG}.{ext}'
def copy(src,dst):
    dst.parent.mkdir(parents=True,exist_ok=True)
    if dst.exists() and Path(src).read_bytes()==dst.read_bytes():return
    shutil.copyfile(src,dst)

# Keep every original project artifact in the archive, including superseded versions.
for directory in ['docs','scripts','images','汇报成果']:
    shutil.copytree(ROOT/directory,ARCH/directory,dirs_exist_ok=True)
copy(ROOT/'README.md',ARCH/'原目录README.md')
copy(ROOT/'.report-build/任务书.txt',ARCH/'课程任务书文字提取件.txt')
req=Path('E:/weichat/xwechat_files/wxid_quoop5hprm8122_9298/msg/file/2026-10/《软件系统分析与设计》课程实践作业课堂展示与提交要求（需求调研与建模成果）.pdf')
copy(req,OUT/'整理依据_课堂展示与提交要求.pdf')

# Existing editable models and exported figures are preserved with clear provenance.
models=B/'03-建模图'
modelmap={}
for n,f in enumerate(sorted((ROOT/'images').glob('*.drawio')),1):
    dst=models/'模型源文件_Drawio'/named(n,f.stem,'drawio');copy(f,dst);modelmap[f.name]=dst
for n,f in enumerate(sorted(list((ROOT/'images').glob('*.png'))+list((ROOT/'scripts').glob('*.png'))),1):
    dst=models/'模型截图'/named(n,f.stem,'png');copy(f,dst);modelmap[f.name]=dst
srs=(ROOT/'docs/系统需求规格说明书_SRS.md').read_text(encoding='utf-8')
for n,code in enumerate(re.findall(r'```mermaid\s*\n(.*?)```',srs,re.S),1):
    title=['系统能力总览','采集任务状态图','门户认领申请状态图'][n-1]
    dst=models/'修订模型源文本_Mermaid'/named(n,title,'mmd');dst.parent.mkdir(parents=True,exist_ok=True);dst.write_text(code,encoding='utf-8')

modelnote='''# 建模材料索引与格式说明

模型源文件 Drawio 目录保存原有可编辑工程，模型截图目录保存原有导出图。修订模型源文本 Mermaid 目录从 v1.1 SRS 第 13 章原样提取能力总览、采集任务状态图和门户认领状态图。

这些文件没有转换为 CodeArts Modeling 工程，不能将 Drawio、Mermaid 或 PNG 声明为 CodeArts 成果。提交前需按课程要求在 CodeArts 中核对并导出模型，补齐源文件。

旧模型截图与修订正文有差异时，按 SRS 第 6.8、6.9、13 章核对状态和授权规则。能力总览采用 Mermaid flowchart，不能代替规范 UML 总用例图。原有活动图部分仅有 PNG，没有对应可编辑源文件；类图和时序图作为辅助材料保留。

文件来源为原目录 images 与 scripts，文件名中的序号仅用于组织材料，不改变原有模型内容。本次未进行模型业务评审或重绘。
'''
(models/'00-模型索引与待补事项.md').write_text(modelnote,encoding='utf-8')

# Research support: retain available records without fabricating interviews.
research=A/'03-调研材料'
research_sources=['学术成果分享平台需求调研报告.md','学术成果分享平台_普通用户需求调研报告.md','科研人员需求分析与需求建模.md','学术成果分享平台_数据源调研与数据密集型逻辑规范.docx','学术成果检索与分享平台_需求调研报告_v0.4.md']
for n,name in enumerate(research_sources,1):copy(ROOT/'scripts'/name,research/'原始调研材料'/named(n,Path(name).stem,Path(name).suffix[1:]))
for n,name in enumerate(['总体方案与需求建模补充.md','学术成果分享平台_详细设计与实现方案汇报.md','15分钟逐页讲稿.md','ppt制作.md'],1):
    copy(ROOT/'汇报成果'/name,B/'05-支撑材料'/'原汇报支撑材料'/named(n,Path(name).stem,'md'))
(B/'05-支撑材料'/'原汇报支撑材料'/'00-使用说明.md').write_text('# 原汇报支撑材料使用说明\n\n此处材料与原 28 页总体汇报关联，讲稿为 20 页主讲共 15 分钟，不能直接作为本次 16 页 PPT 的逐页讲稿。其日期、需求编号和状态规则可能早于 v1.1，以提交包正文为准。课程建议汇报 12 分钟、问答 3 分钟，需另行排练和调整。\n',encoding='utf-8')

revision='''# 修改说明与版本记录

本表记录 2026 年 10 月 12 日按课堂展示与提交要求进行的格式整理。正文沿用原目录 docs 三份 v1.1 稿件，原稿记载日期为 2026 年 10 月 10 日。本次未收到教师或互评意见，未填写虚构反馈、修改人或人工验证结论。

组号：待填。团队名：待填。提交负责人：待填。

## 本次整理记录

| 序号 | 修改项 | 反馈来源 教师或互评 | 修改内容 | 修改人 | 修改日期 | 验证方式 |
| --- | --- | --- | --- | --- | --- | --- |
| 1 | 交付组织 | 无 本次按提交要求整理 | 分开组织作业 1 与作业 2 并生成各自压缩包 | AI 辅助整理 人工负责人待填 | 2026-10-12 | 对照材料清单及压缩包目录 |
| 2 | 正文格式 | 无 本次按提交要求整理 | 将 v1.1 调研报告 SRS SPEC 转为 Word 统一 A4 字体 标题 表格和页码 | AI 辅助整理 人工负责人待填 | 2026-10-12 | 正文内容核对及排版检查 |
| 3 | 文档引用 | 无 本次整理发现 | 修复提交副本失效相对路径 补建本次整合说明 保留原稿业务表述 | AI 辅助整理 人工负责人待填 | 2026-10-12 | 本地引用目标存在性检查 |
| 4 | 模型附件 | 无 本次按提交要求整理 | 分类复制 Drawio 源文件和 PNG 从 SRS 提取 Mermaid 源文本 标明待补 CodeArts 工程 | AI 辅助整理 人工负责人待填 | 2026-10-12 | 文件清单及来源核对 |
| 5 | 通用附件 | 无 本次按提交要求整理 | 按附录 C 和 D 建立修改记录与自查清单 补建 AI 使用声明 | AI 辅助整理 人工负责人待填 | 2026-10-12 | 与要求表头和材料项逐项比对 |

## 课堂反馈记录 待课堂后填写

| 序号 | 修改项 | 反馈来源 教师或互评 | 修改内容 | 修改人 | 修改日期 | 验证方式 |
| --- | --- | --- | --- | --- | --- | --- |
| 1 | 待填写 | 待填写 | 待填写 | 待填写 | 待填写 | 待填写 |
| 2 | 待填写 | 待填写 | 待填写 | 待填写 | 待填写 | 待填写 |

教师与互评提出的事实性错误和交付遗漏需要修改；建议项应记录采纳情况和理由。课程展示尚未进行，本表不能作为展示后反馈已闭环的证明。
'''
ai='''# AI 使用声明

本声明记录本次材料整理中可确认的 AI 使用情况。历史稿件使用了哪些模型、提示词及人工验证步骤，当前目录不能完整确认，需要团队补充。

## 本次使用情况

使用工具为 Codex 中的 AI 助手，时间为 2026 年 10 月 12 日。AI 读取用户提供的课堂展示与提交要求和本地项目材料，完成文件分类、命名、Markdown 到 Word 的排版转换、相对链接修复，以及修改记录、自查清单和本声明的草稿编写。从 SRS 提取的 Mermaid 源文本沿用原稿。

AI 未开展真实访谈或问卷调查，未重新执行原报告所记载的数据采集实验，未进行课程评审，未替团队确认模型已符合 CodeArts 要求。原稿中的实验和性能目标保留其待核验状态。

## 验证情况与团队责任

本次自动核对包括文件来源、正文转换、附件清单和本地引用路径。自动核对不能代替团队对需求、模型关系、算法规则、外部资料、数据合规和课程要求的理解与验证。团队人工审查与签字尚未完成。

提交前，团队需逐项审查正文和建模图，填写修改记录与自查表，补齐真实证据，并明确 AI 生成或整理内容的使用范围。最终判断与提交责任由团队承担。

## 历史 AI 使用记录 待团队补充

| 环节 | 模型或工具与版本 | 提示词和生成内容 | 人工验证方式 | 验证人和日期 |
| --- | --- | --- | --- | --- |
| 需求调研 | 待填 | 待填 | 待填 | 待填 |
| 需求建模 | 待填 | 待填 | 待填 | 待填 |
| 两个小组工具路线比较 | 待填 | 待填 | 待填 | 待填 |

团队名称：待填。成员及分工：待填。人工审查人：待填。签字或确认日期：待填。
'''
check='''# 提交材料自查清单

本清单按要求附录 D 整理，记录现有文件状态。目录整理完成不等于课程终稿已全部通过。两份作业包分别带有本清单，便于提交负责人对照同一份材料状态核验。

| 序号 | 材料 | 是否完成 | 自查人 | 备注 |
| --- | --- | --- | --- | --- |
| 1 | 需求调研报告 Word | 已整理 待人工审查 | 待填 | 作业 1 中的 v1.1 Word 正文 |
| 2 | 需求调研汇报 PPT | 已提供 待修订确认 | 待填 | 复制现有综合汇报 软分展示2 共 16 页 未拆分内容 课程建议 10 至 15 页 |
| 3 | 调研材料 | 部分具备 | 待填 | 已有角色分析 数据源和平台比较记录 无真实访谈问卷记录 原稿两项实验原件待补 |
| 4 | 系统需求规格说明书 SRS Word | 已整理 待人工审查 | 待填 | 作业 2 中的 v1.1 Word 正文 |
| 5 | 核心模块开发规范 SPEC Word | 已整理 待人工审查 | 待填 | 已保留清洗 去重 消歧 检索排序规则 实测与评审另行完成 |
| 6 | 用例图 CodeArts 源文件及导出图 | 未全部完成 | 待填 | 已有 Drawio 和 PNG 尚无 CodeArts 工程 需与 v1.1 新增功能同步 |
| 7 | 活动图源文件及导出图 | 未全部完成 | 待填 | 有多张 PNG 部分可编辑源文件缺失 需核对并在 CodeArts 中更新 |
| 8 | 状态图源文件及导出图 | 未全部完成 | 待填 | 已有历史 Drawio PNG 及从 SRS 提取的修订 Mermaid 文本 需更新导出图和 CodeArts 工程 |
| 9 | 需求建模汇报 PPT 与支撑材料 | 已提供 待修订确认 | 待填 | 同一 16 页综合汇报放入两份作业 原 28 页讲稿另作参考 当前 PPT 封面缺成员分工 |
| 10 | 修改说明与版本记录表 | 已建表 待课堂后补充 | 待填 | 已记录本次整理 无教师互评意见 需课堂后追加修订 |
| 11 | AI 使用声明 | 已建稿 待人工确认 | 待填 | 本次使用范围已写明 历史模型提示词 两小组比较和人工签字待补 |

## 提交前需要完成

填写组号 团队名 成员分工和提交负责人，替换文件及压缩包中的待填字样。根据课堂展示版本核对 PPT，补全封面和 AI 使用说明，并按 12 分钟汇报加 3 分钟问答排练。同步 SRS 和模型，补齐 CodeArts 工程及导出图，确认缺失证据并按实际情况填写。课堂后补教师和互评意见以及人工验证记录。全部项目通过后由一名指定负责人分作业提交。
'''
integration='''# 文档整合说明

本说明新建于 2026 年 10 月 12 日，用于提交包的格式与文件组织。原目录引用的旧版整合说明未找到，不能据此恢复未记录的历史修改。

## 正文依据

正式正文选自 docs 中的需求调研报告、系统需求规格说明书 SRS 和核心模块开发规范 SPEC，三份均自标 v1.1 和 2026 年 10 月 10 日。Word 是本次对这些正文的排版转换，业务描述和不确定性均予保留。文本源稿中的相对链接已适配提交包目录。

## 支撑与历史材料

角色调研和数据源报告作为调研支撑；Drawio 源文件、原 PNG 和从 SRS 提取的 Mermaid 文本作为模型附件。各源图可能早于 v1.1，需人工核对后在课程指定工具中更新。现有综合演示文稿 软分展示2 被原样复制到两份作业包，作为含调研与建模内容的共同汇报。它包含 16 页，未作内容重写。其他 PPT 和历史 Word 统一放入 Submit 根目录的历史归档，不随两个压缩包提交。

## 待补事项

组号、团队名、成员分工、课堂反馈、人工审查记录、历史 AI 使用和两个小组模型路线比较尚未提供。原稿声明真实访谈或问卷未开展，E1 E2 实验脚本、原始响应和样本未随材料提交。课程任务书 PDF、早期 SRS v0.1 及部分模型说明也仍待补；已保留目录中存在的任务书文字提取件。此处不填造缺失记录。
'''

# All links inside each independently submitted package remain local to that package.
unique_docs=[]
mdnames={name:named(i,Path(name).stem,'md') for i,name in enumerate(['需求调研报告.md','系统需求规格说明书_SRS.md','核心模块开发规范_SPEC.md'],1)}
for package in [A,B]:
    texts=package/'文本源稿';texts.mkdir(exist_ok=True)
    note=texts/named(4,'整合说明','md');note.write_text(integration,encoding='utf-8')
    aux=package/'参考附件'
    copy(ROOT/'scripts/学术成果检索与分享平台_核心模块开发规范_SPEC_v0.1_会前评审稿.md',aux/named(1,'SPEC会前评审稿','md'))
    copy(ROOT/'.report-build/任务书.txt',aux/named(2,'课程任务书文字提取件','txt'))
    # Both packages need the figures referenced by their Markdown originals.
    if package==A:
        for src,dst in modelmap.items():copy(dst,aux/dst.name)
    for name,newname in mdnames.items():
        source=(ROOT/'docs'/name).read_text(encoding='utf-8')
        def link(m):
            target=m.group(2)
            if re.match(r'^(https?://|#|mailto:)',target):return m.group(0)
            base=Path(target.split('#')[0]).name
            if base in mdnames:dest=texts/mdnames[base]
            elif base=='整合说明.md':dest=note
            elif base in modelmap:dest=(aux/modelmap[base].name) if package==A else modelmap[base]
            elif base=='学术成果检索与分享平台_核心模块开发规范_SPEC_v0.1_会前评审稿.md':dest=aux/named(1,'SPEC会前评审稿','md')
            else:
                found=[p for folder in ['scripts','docs','汇报成果','images'] for p in (ROOT/folder).glob(base)]
                if not found:raise ValueError(f'Unmapped link {name}: {target}')
                dest=aux/('90-'+base);copy(found[0],dest)
            return f'{m.group(1)}({Path(os.path.relpath(dest,texts)).as_posix()})'
        source=re.sub(r'(!?\[[^\]]*\])\(([^)]+)\)',link,source)
        (texts/newname).write_text(source,encoding='utf-8')
    for n,title,content in [(4 if package==A else 6,'修改说明与版本记录',revision),(5 if package==A else 7,'AI使用声明',ai),(6 if package==A else 8,'提交材料自查清单',check)]:
        src=package/named(n,title,'md');src.write_text(content,encoding='utf-8')
        if package==A:unique_docs.append((src,src.with_suffix('.docx')))
    copy(ROOT/'汇报成果/软分展示2.pptx',package/named(2 if package==A else 4,'需求调研汇报' if package==A else '需求建模汇报','pptx'))

# A readable extract of actually available research evidence, explicitly marked as excerpts.
research_text=(ROOT/'docs/需求调研报告.md').read_text(encoding='utf-8')
sections=re.split(r'(?=^## )',research_text,flags=re.M)
extract='# 调研支撑材料摘编\n\n本材料从 v1.1 需求调研报告摘录调研方法、三类角色、平台比较、数据来源与资料索引，便于单独核查。内容仍为原报告的公开资料分析与场景推演，未新增访谈或实验原件。完整背景、结论和限制请结合正式报告阅读。\n\n'
extract+='\n\n'.join(s for s in sections if re.match(r'## (1\.|3\.|5\.|8\.|12\.)',s))
# Use the already corrected package version to keep excerpt links valid.
corrected=(A/'文本源稿'/mdnames['需求调研报告.md']).read_text(encoding='utf-8')
sections=re.split(r'(?=^## )',corrected,flags=re.M)
extract=extract.split('\n\n## ')[0]+'\n\n'+'\n\n'.join(s for s in sections if re.match(r'## (1\.|3\.|5\.|8\.|12\.)',s))
extract=re.sub(r'(!?\[[^\]]*\])\(([^)]+)\)',lambda m:m.group(0) if re.match(r'^(https?://|#)',m.group(2)) else f'{m.group(1)}(../文本源稿/{m.group(2)})',extract)
extract_path=research/named(6,'调研支撑材料摘编','md');extract_path.write_text(extract,encoding='utf-8');unique_docs.append((extract_path,extract_path.with_suffix('.docx')))
for name,pkg,n,title in [('需求调研报告.md',A,1,'需求调研报告'),('系统需求规格说明书_SRS.md',B,1,'SRS'),('核心模块开发规范_SPEC.md',B,2,'SPEC')]:
    unique_docs.append((pkg/'文本源稿'/mdnames[name],pkg/named(n,title,'docx')))

def setfont(style,font,size,bold=False):
    style.font.name=font;style.font.size=Pt(size);style.font.bold=bold;style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().append(OxmlElement('w:rFonts'))
    style.element.rPr.rFonts.set(qn('w:eastAsia'),font)

ref=Document()
setfont(ref.styles['Normal'],'宋体',10.5)
for style,font,size in [('Title','黑体',20),('Subtitle','宋体',11),('Heading 1','黑体',16),('Heading 2','黑体',13),('Heading 3','黑体',11),('Heading 4','黑体',10.5)]:setfont(ref.styles[style],font,size,style.startswith('Heading'))
setfont(ref.styles['Source Code'] if 'Source Code' in ref.styles else ref.styles.add_style('Source Code',1),'Consolas',8)
ref.save(BUILD/'reference.docx')

def format_doc(path):
    d=Document(path)
    for sec in d.sections:
        sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(2.1);sec.bottom_margin=Cm(2.0);sec.left_margin=Cm(2.0);sec.right_margin=Cm(2.0)
        sec.header_distance=Cm(.8);sec.footer_distance=Cm(.8)
        p=sec.header.paragraphs[0];p.text='软件系统分析与设计  学术成果检索与分享平台';p.style='Normal'
        for r in p.runs:r.font.size=Pt(8)
        p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.add_run('第 ');f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f);p.add_run(' 页')
    for p in d.paragraphs:
        pf=p.paragraph_format;pf.space_after=Pt(5);pf.line_spacing=1.2;pf.widow_control=True
        if p.style.name.startswith('Heading'):
            pf.keep_with_next=True;pf.space_before=Pt(10);pf.space_after=Pt(5)
        if p.style.name=='Title':
            pf.space_before=Pt(0);pf.space_after=Pt(12)
            for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
        if p.style.name=='Source Code':
            pf.line_spacing=1.05;pf.space_after=Pt(1);pf.keep_with_next=False
            for r in p.runs:
                r.font.size=Pt(8);r.font.name='Consolas';r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
        for r in p.runs:
            for inline in r._r.xpath('.//wp:inline'):
                ext=inline.find(qn('wp:extent'))
                if ext is not None:
                    cx=int(ext.get('cx'));cy=int(ext.get('cy'));ratio=min(Cm(16.7)/cx,Cm(19)/cy,1)
                    ext.set('cx',str(int(cx*ratio)));ext.set('cy',str(int(cy*ratio)))
                    for x in inline.xpath('.//a:xfrm/a:ext'):x.set('cx',str(int(cx*ratio)));x.set('cy',str(int(cy*ratio)))
    for t in d.tables:
        t.autofit=False
        cols=len(t.columns)
        lens=[max(len(row.cells[i].text) for row in t.rows) for i in range(cols)]
        weights=[max(4,min(100,n))**.55 for n in lens];widths=[17*w/sum(weights) for w in weights]
        for i,col in enumerate(t.columns):col.width=Cm(widths[i])
        borders=OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
        t._tbl.tblPr.append(borders)
        for ri,row in enumerate(t.rows):
            pr=row._tr.get_or_add_trPr()
            if ri==0:
                e=OxmlElement('w:tblHeader');pr.append(e)
            # Allow long narrative rows to split instead of generating large page gaps.
            for ci,c in enumerate(row.cells):
                c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcp=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
                for side in ['top','left','bottom','right']:
                    e=OxmlElement('w:'+side);e.set(qn('w:w'),'75');e.set(qn('w:type'),'dxa');mar.append(e)
                tcp.append(mar)
                if ri==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E5EAF0');tcp.append(shade)
                for p in c.paragraphs:
                    p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1);p.paragraph_format.line_spacing=1.12;p.paragraph_format.keep_with_next=False;p.paragraph_format.widow_control=True
                    for r in p.runs:
                        r.font.size=Pt(9);r.font.color.rgb=RGBColor(0,0,0);r.font.name='宋体';r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
                        if ri==0:r.bold=True
        d.add_paragraph() if t==d.tables[-1] else None
    d.core_properties.author='课程项目团队 待填写';d.core_properties.subject='提交材料格式整理';d.core_properties.comments='源正文 v1.1；整理日期 2026-10-12；人工审核待完成'
    d.save(path)

for source,dest in unique_docs:
    text=source.read_text(encoding='utf-8')
    first=text.splitlines()[0].lstrip('# ').strip()
    body='\n'.join(text.splitlines()[1:])
    draft=BUILD/(dest.stem+'.md');draft.write_text(body,encoding='utf-8')
    subprocess.run(['pandoc',str(draft),'-f','markdown','-t','docx','--metadata',f'title={first}','--resource-path',str(source.parent),'--reference-doc',str(BUILD/'reference.docx'),'-o',str(dest)],check=True,cwd=source.parent)
    format_doc(dest)
    print('CREATED',dest.relative_to(OUT))
for title,n1,n2 in [('修改说明与版本记录',4,6),('AI使用声明',5,7),('提交材料自查清单',6,8)]:copy(A/named(n1,title,'docx'),B/named(n2,title,'docx'))

readme='''# Submit 提交材料说明

本目录按课程要求分成两份作业，保留原目录全部文件。三份 v1.1 正文已排版为 Word；现有综合 PPT 分别放入两份作业包；模型按源文件、导出图和修订源文本分类。根目录历史归档仅供查阅，不放入作业压缩包。

| 目录或文件 | 用途 |
| --- | --- |
| 作业1_需求调研成果 | 需求调研报告 Word 汇报 PPT 调研支撑材料与通用附件 |
| 作业2_需求建模成果 | SRS Word SPEC Word 建模附件 汇报 PPT 与通用附件 |
| 两份 zip | 分别对应作业 1 和作业 2 当前为待完善材料包 |
| 归档_历史材料_不放入提交压缩包 | 所有原始文件 按原来源保存 不作为最新提交正文 |
| 整理依据_课堂展示与提交要求.pdf | 用户提供的课程要求副本 |
| 文件清单.csv | 文件路径 大小 内容校验值及压缩包归属 |

## 提交前待办

当前没有组号和团队名，文件使用“组号待填-团队名待填”，不可原样作为最终命名提交。替换名称后需重新打包。团队成员、分工、历史 AI 使用及人工验证情况也需填写。

现有 软分展示2 综合汇报共 16 页，被原样复制到两份作业包；课程建议 10 至 15 页，封面成员分工、AI 声明和 12 分钟讲述需人工完善。原 28 页汇报和不同 16 页版本已留在历史归档，防止误把多个版本作为正式稿。支撑材料中的原讲稿仅对应 28 页版本。

现有模型为 Drawio、PNG 及 SRS 中提取的 Mermaid 文本，未转换成 CodeArts 工程。旧截图与修订 SRS 的生命周期可能不同，需在 CodeArts 中同步并补齐源文件和导出图。真实访谈问卷、E1 E2 实验原件、课堂反馈和人工确认记录仍待补。各自查清单如实保留“待审查”“部分具备”“未全部完成”状态。

课堂展示后，补齐教师与互评意见及修订验证记录，全部自查通过后再提交终稿。课程文件建议 2026 年 10 月 20 日 24:00 前提交，最终以教师通知为准。本次仅整理本地文件，未向教学平台或教师群提交。
'''
(OUT/'README.md').write_text(readme,encoding='utf-8')
(BUILD/'docx_outputs.json').write_text(json.dumps([str(d) for _,d in unique_docs],ensure_ascii=False,indent=2),encoding='utf-8')
print('READY FOR RENDER',len(unique_docs))
